"""Atualiza a planilha de carga (ACWR) preservando fórmulas e named ranges.

Escreve apenas as colunas de entrada:

- ``Daily_Data``: B (Data), C (Carga diária Km)
- ``PACE``: B (Data), C (Carga Km), D (PACE), E (Tempo total)

Mapeia data -> linha a partir da data de início (`Dia 1` = `start_date` do
corredor), usando `Corredor.dia_da_planilha()`:

    dia = (data - start_date).days + 1  ->  linha = dia + 1

**`B2` de `Daily_Data` é a fonte da verdade** depois da primeira escrita: se o
`corredores.toml` divergir do que já está gravado, a sincronização do corredor
é abortada (`DataBaseDivergenteError`) em vez de reescrever a grade deslocada.
Planilha vazia (`B2` em branco) é o caso de bootstrap — a primeira gravação
preenche `B2` naturalmente, ao escrever o `Dia 1`.

A grade tem teto (`Daily_Data` até a linha 366 — 365 dias; `PACE` até a 154 —
153 dias, ~5 meses). Um dia além do limite de uma aba é **descartado só
naquela aba**, com um WARNING no log — não interrompe a sincronização do
corredor nem a gravação nas demais linhas/abas. O dado de pesquisa continua
íntegro no SQLite; decidir se/quando abrir uma segunda planilha por corredor
fica para quando o problema aparecer de verdade.
"""
from __future__ import annotations

import os
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import Cell
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.models.corredor import Corredor
from src.models.daily_load import DailyLoad
from src.utils.errors import (
    DataBaseDivergenteError,
    PlanilhaEmUsoError,
    PlanilhaError,
    PlanilhaNaoEncontradaError,
)
from src.utils.logger import get_logger

logger = get_logger(__name__)

ABA_DAILY_DATA = "Daily_Data"
ABA_PACE = "PACE"

# Linha 1 é cabeçalho; `Dia 1` mora na linha 2 (ver `Corredor.dia_da_planilha`).
PRIMEIRA_LINHA = 2
ULTIMA_LINHA_DAILY_DATA = 366  # 365 dias
ULTIMA_LINHA_PACE = 154  # 153 dias

# Colunas de entrada — as únicas que este serviço toca. D-J de Daily_Data são
# fórmulas (EWMA, ACWR, ...) e nunca aparecem aqui.
COL_DATA = 2  # B
COL_CARGA = 3  # C
COL_PACE = 4  # D — só em PACE
COL_TEMPO = 5  # E — só em PACE

FORMATO_DATA = "yyyy-mm-dd"
FORMATO_CARGA = "0.00"
FORMATO_PACE = "0.00"
FORMATO_TEMPO = "[h]:mm:ss"


@dataclass(frozen=True)
class ResultadoExcel:
    """O que uma sincronização escreveu na planilha de um corredor."""

    dias_escritos_daily_data: int
    dias_escritos_pace: int
    # Anteriores ao Dia 1 — fora do escopo do estudo, não é erro de ninguém.
    dias_descartados: int = 0
    # Depois do teto de cada aba — dado de pesquisa, mas sem linha para ir.
    dias_alem_do_limite_daily_data: int = 0
    dias_alem_do_limite_pace: int = 0

    @property
    def dias_ignorados(self) -> int:
        return (
            self.dias_descartados
            + self.dias_alem_do_limite_daily_data
            + self.dias_alem_do_limite_pace
        )


class ExcelService:
    """Escreve o agregado diário nas abas Daily_Data e PACE, sem tocar fórmulas."""

    def sincronizar(self, corredor: Corredor, daily_loads: Iterable[DailyLoad]) -> ResultadoExcel:
        """Grava os `DailyLoad` na planilha do corredor.

        Raises:
            PlanilhaNaoEncontradaError: o corredor ainda não fez o onboarding
                (o `excel_path` não existe).
            PlanilhaEmUsoError: o arquivo está aberto em outro programa.
            DataBaseDivergenteError: `B2` já gravado diverge do `start_date`
                do cadastro.
            PlanilhaError: qualquer outra falha de leitura/escrita do arquivo.
        """
        loads = sorted(daily_loads, key=lambda carga: carga.day)
        if not loads:
            logger.debug("Corredor %s: nenhum dia para escrever no Excel.", corredor)
            return ResultadoExcel(0, 0)

        workbook = self._abrir(corredor)
        aba_daily = workbook[ABA_DAILY_DATA]
        aba_pace = workbook[ABA_PACE]

        self._validar_data_base(corredor, aba_daily)

        descartados = 0
        alem_daily = 0
        alem_pace = 0
        escritos_daily = 0
        escritos_pace = 0

        for carga in loads:
            dia = corredor.dia_da_planilha(carga.day)
            if dia <= 0:
                logger.warning(
                    "Corredor %s: %s é anterior ao Dia 1 (%s) — fora do escopo "
                    "do estudo, não escrito.",
                    corredor,
                    carga.day.isoformat(),
                    corredor.start_date.isoformat(),
                )
                descartados += 1
                continue

            linha = dia + 1

            if _escrever_daily_data(aba_daily, linha, carga):
                escritos_daily += 1
            else:
                logger.warning(
                    "Corredor %s: %s (dia %d) além do limite de %s (linha %d) "
                    "— não escrito nessa aba.",
                    corredor,
                    carga.day.isoformat(),
                    dia,
                    ABA_DAILY_DATA,
                    ULTIMA_LINHA_DAILY_DATA,
                )
                alem_daily += 1

            if _escrever_pace(aba_pace, linha, carga):
                escritos_pace += 1
            else:
                logger.warning(
                    "Corredor %s: %s (dia %d) além do limite de %s (linha %d) "
                    "— não escrito nessa aba.",
                    corredor,
                    carga.day.isoformat(),
                    dia,
                    ABA_PACE,
                    ULTIMA_LINHA_PACE,
                )
                alem_pace += 1

        resultado = ResultadoExcel(
            dias_escritos_daily_data=escritos_daily,
            dias_escritos_pace=escritos_pace,
            dias_descartados=descartados,
            dias_alem_do_limite_daily_data=alem_daily,
            dias_alem_do_limite_pace=alem_pace,
        )

        if escritos_daily == 0 and escritos_pace == 0:
            # Nada mudou no conteúdo — não vale abrir mão do arquivo original
            # (e do risco de "planilha aberta no Excel") por uma gravação sem
            # efeito nenhum.
            logger.info("Corredor %s: nada gravado no Excel (%s).", corredor, resultado)
            return resultado

        self._salvar(corredor, workbook)
        logger.info("Corredor %s: planilha atualizada (%s).", corredor, resultado)
        return resultado

    # --------------------------------------------------------------- interno

    def _abrir(self, corredor: Corredor) -> Workbook:
        caminho = corredor.excel_path
        if not caminho.is_file():
            raise PlanilhaNaoEncontradaError(
                f"corredor {corredor.id}: planilha não encontrada em {caminho} "
                "(copie o template para este caminho antes de sincronizar)"
            )
        try:
            return openpyxl.load_workbook(caminho)
        except PermissionError as erro:
            raise PlanilhaEmUsoError(
                f"corredor {corredor.id}: não foi possível abrir {caminho.name} "
                "(o arquivo parece estar aberto em outro programa; feche-o e "
                "rode a sincronização de novo)"
            ) from erro
        except (OSError, InvalidFileException) as erro:
            raise PlanilhaError(
                f"corredor {corredor.id}: falha ao abrir {caminho}: {erro}"
            ) from erro

    def _validar_data_base(self, corredor: Corredor, aba_daily: Worksheet) -> None:
        """`B2` em branco é bootstrap (planilha nova); gravado, é a fonte da verdade."""
        bruto = aba_daily.cell(row=PRIMEIRA_LINHA, column=COL_DATA).value
        if bruto is None:
            return

        gravado = _como_data(bruto)
        if gravado is None:
            raise DataBaseDivergenteError(
                f"corredor {corredor.id}: {ABA_DAILY_DATA}!B2 não é uma data "
                f"reconhecível ({bruto!r}) — confira a planilha antes de sincronizar"
            )

        if gravado != corredor.start_date:
            raise DataBaseDivergenteError(
                f"corredor {corredor.id}: {ABA_DAILY_DATA}!B2={gravado.isoformat()} "
                f"diverge do start_date={corredor.start_date.isoformat()} no cadastro "
                "— a grade já foi gravada com outra data-base. Corrija o "
                "corredores.toml para bater com a planilha; mudar B2 deslocaria "
                "todos os dias já gravados."
            )

    def _salvar(self, corredor: Corredor, workbook: Workbook) -> None:
        """Grava num arquivo temporário e troca com `os.replace` (atômico).

        Uma falha no meio do `wb.save()` não pode deixar a planilha do
        corredor — o produto da pesquisa — corrompida pela metade.
        """
        destino = corredor.excel_path
        temporario = destino.with_name(f".{destino.stem}.tmp{destino.suffix}")

        try:
            workbook.save(temporario)
        except OSError as erro:
            _apagar(temporario)
            raise PlanilhaError(
                f"corredor {corredor.id}: falha ao gravar {destino}: {erro}"
            ) from erro

        try:
            os.replace(temporario, destino)
        except PermissionError as erro:
            _apagar(temporario)
            raise PlanilhaEmUsoError(
                f"corredor {corredor.id}: não foi possível salvar em {destino.name} "
                "(o arquivo parece estar aberto em outro programa; feche-o e "
                "rode a sincronização de novo)"
            ) from erro
        except OSError as erro:
            _apagar(temporario)
            raise PlanilhaError(
                f"corredor {corredor.id}: falha ao substituir {destino}: {erro}"
            ) from erro


def _escrever_daily_data(aba: Worksheet, linha: int, carga: DailyLoad) -> bool:
    """Escreve B (Data) e C (Carga). Devolve False se `linha` passou do teto."""
    if linha > ULTIMA_LINHA_DAILY_DATA:
        return False
    _escrever_data(aba.cell(row=linha, column=COL_DATA), carga.day)
    _escrever_numero(aba.cell(row=linha, column=COL_CARGA), carga.carga_km, FORMATO_CARGA)
    return True


def _escrever_pace(aba: Worksheet, linha: int, carga: DailyLoad) -> bool:
    """Escreve B/C/D/E. Devolve False se `linha` passou do teto.

    Pace e Tempo ficam em branco em dia de descanso — só a Carga é 0, para que
    a grade continue contígua sem sugerir um pace de "zero minutos por km".
    """
    if linha > ULTIMA_LINHA_PACE:
        return False
    _escrever_data(aba.cell(row=linha, column=COL_DATA), carga.day)
    _escrever_numero(aba.cell(row=linha, column=COL_CARGA), carga.carga_km, FORMATO_CARGA)
    _escrever_numero(aba.cell(row=linha, column=COL_PACE), carga.pace_min_km, FORMATO_PACE)
    _escrever_tempo(aba.cell(row=linha, column=COL_TEMPO), carga.tempo_total)
    return True


def _escrever_data(celula: Cell, valor: date) -> None:
    celula.value = valor
    celula.number_format = FORMATO_DATA


def _escrever_numero(celula: Cell, valor: float | None, formato: str) -> None:
    celula.value = valor
    if valor is not None:
        celula.number_format = formato


def _escrever_tempo(celula: Cell, valor: timedelta | None) -> None:
    celula.value = valor
    if valor is not None:
        celula.number_format = FORMATO_TEMPO


def _como_data(valor: object) -> date | None:
    """Lê o que está em `B2` como `date`, aceitando data nativa ou texto ISO."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    try:
        return date.fromisoformat(str(valor).strip())
    except ValueError:
        return None


def _apagar(caminho: Path) -> None:
    """Remove um arquivo temporário, sem levantar se ele já sumiu."""
    try:
        caminho.unlink(missing_ok=True)
    except OSError:
        logger.debug("Não foi possível remover o temporário %s.", caminho)
