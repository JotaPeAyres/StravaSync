"""Testes do AdocaoService (Fase 5.5): ExcelService + ActivityRepository juntos.

O contrato importante é o que o resto do projeto já garante em separado —
`ExcelService.ler_historico_manual` e `ActivityRepository.importar_historico`
têm testes próprios — este arquivo verifica só a costura entre os dois e o
relatório que sai no fim.
"""
from __future__ import annotations

import shutil
from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pytest

from src.models.corredor import Corredor
from src.repositories.activity_repository import ActivityRepository
from src.services.adocao_service import AdocaoService, ResultadoAdocao
from src.services.database_service import DatabaseService
from src.services.excel_service import ABA_DAILY_DATA, ABA_PACE, COL_CARGA, COL_DATA, COL_TEMPO
from src.utils.errors import (
    GradeDesalinhadaError,
    PlanilhaEmUsoError,
    PlanilhaNaoEncontradaError,
)

TEMPLATE = Path(__file__).resolve().parents[1] / "Cópia de Planilha_carga_corrida.xlsx"
INICIO = date(2026, 1, 1)

pytestmark = pytest.mark.skipif(
    not TEMPLATE.is_file(), reason="template real não encontrado no repositório"
)


@pytest.fixture
def planilha(tmp_path) -> Path:
    destino = tmp_path / "p001.xlsx"
    shutil.copy(TEMPLATE, destino)
    return destino


@pytest.fixture
def banco(tmp_path):
    servico = DatabaseService(tmp_path / "estado.db")
    servico.init_schema()
    try:
        yield servico
    finally:
        servico.close()


@pytest.fixture
def servico(banco):
    from src.services.excel_service import ExcelService

    return AdocaoService(ExcelService(), ActivityRepository(banco.connect()))


def _corredor(excel_path: Path, *, cutover_date: date) -> Corredor:
    return Corredor(
        id="p001",
        nome="Ana",
        refresh_token="token",
        start_date=INICIO,
        cutover_date=cutover_date,
        excel_path=excel_path,
    )


def _preencher(wb, corredor: Corredor, dia: date, carga_km: float, minutos: float | None = None):
    linha = corredor.dia_da_planilha(dia) + 1
    wb[ABA_DAILY_DATA].cell(row=linha, column=COL_DATA).value = dia
    wb[ABA_DAILY_DATA].cell(row=linha, column=COL_CARGA).value = carga_km
    if minutos is not None:
        wb[ABA_PACE].cell(row=linha, column=COL_DATA).value = dia
        wb[ABA_PACE].cell(row=linha, column=COL_TEMPO).value = timedelta(minutes=minutos)


def test_corredor_novo_nao_adota_nada(planilha, servico):
    corredor = _corredor(planilha, cutover_date=INICIO)  # start == cutover

    resultado = servico.adotar(corredor)

    assert resultado == ResultadoAdocao(
        corredor_id="p001",
        dias_importados=0,
        dias_sem_corrida=0,
        lacunas=(),
        dias_sem_tempo_conhecido=(),
        primeiro_dia_sob_gestao_do_app=INICIO,
    )


def test_adota_periodo_manual_completo(planilha, banco, servico):
    cutover = INICIO + timedelta(days=3)
    corredor = _corredor(planilha, cutover_date=cutover)

    wb = openpyxl.load_workbook(planilha)
    _preencher(wb, corredor, INICIO, 10.0, 50.0)
    _preencher(wb, corredor, INICIO + timedelta(days=1), 0.0)  # descanso
    _preencher(wb, corredor, INICIO + timedelta(days=2), 8.0, 45.0)
    wb.save(planilha)

    resultado = servico.adotar(corredor)

    assert resultado.dias_importados == 2  # só as corridas de fato
    assert resultado.dias_sem_corrida == 1
    assert resultado.lacunas == ()
    assert resultado.dias_sem_tempo_conhecido == ()
    assert resultado.primeiro_dia_sob_gestao_do_app == cutover

    repo = ActivityRepository(banco.connect())
    achadas = repo.por_periodo("p001", INICIO, cutover - timedelta(days=1))
    assert len(achadas) == 2
    assert all(a.id is None for a in achadas)


def test_relata_lacunas_e_dias_sem_tempo(planilha, servico):
    cutover = INICIO + timedelta(days=3)
    corredor = _corredor(planilha, cutover_date=cutover)

    wb = openpyxl.load_workbook(planilha)
    _preencher(wb, corredor, INICIO, 10.0, 50.0)
    # INICIO + 1 dia fica em branco (lacuna).
    _preencher(wb, corredor, INICIO + timedelta(days=2), 8.0)  # sem PACE: tempo desconhecido
    wb.save(planilha)

    resultado = servico.adotar(corredor)

    assert resultado.lacunas == (INICIO + timedelta(days=1),)
    assert resultado.dias_sem_tempo_conhecido == (INICIO + timedelta(days=2),)
    assert resultado.tem_pendencias is True


def test_adocao_e_idempotente(planilha, banco, servico):
    cutover = INICIO + timedelta(days=2)
    corredor = _corredor(planilha, cutover_date=cutover)

    wb = openpyxl.load_workbook(planilha)
    _preencher(wb, corredor, INICIO, 10.0, 50.0)
    wb.save(planilha)

    servico.adotar(corredor)
    servico.adotar(corredor)

    repo = ActivityRepository(banco.connect())
    achadas = repo.por_periodo("p001", INICIO, INICIO)
    assert len(achadas) == 1  # não duplicou


def test_nao_escreve_no_excel(planilha, servico):
    """Adotar só lê a planilha e grava no SQLite — o Excel já tem esse histórico."""
    cutover = INICIO + timedelta(days=2)
    corredor = _corredor(planilha, cutover_date=cutover)

    wb = openpyxl.load_workbook(planilha)
    _preencher(wb, corredor, INICIO, 10.0, 50.0)
    wb.save(planilha)
    mtime_antes = planilha.stat().st_mtime_ns

    servico.adotar(corredor)

    assert planilha.stat().st_mtime_ns == mtime_antes


def test_relatorio_sem_pendencias_quando_tudo_completo(planilha, servico):
    cutover = INICIO + timedelta(days=2)
    corredor = _corredor(planilha, cutover_date=cutover)

    wb = openpyxl.load_workbook(planilha)
    _preencher(wb, corredor, INICIO, 10.0, 50.0)
    _preencher(wb, corredor, INICIO + timedelta(days=1), 0.0)
    wb.save(planilha)

    resultado = servico.adotar(corredor)

    assert resultado.tem_pendencias is False


# --------------------------------------------------------------------- erros


def test_grade_desalinhada_propaga_sem_gravar_nada_parcial(planilha, banco, servico):
    """Linha com data errada é sinal de linha inserida/apagada à mão (Fase 5.5)."""
    cutover = INICIO + timedelta(days=3)
    corredor = _corredor(planilha, cutover_date=cutover)

    wb = openpyxl.load_workbook(planilha)
    _preencher(wb, corredor, INICIO, 10.0, 50.0)
    linha_errada = corredor.dia_da_planilha(INICIO + timedelta(days=1)) + 1
    wb[ABA_DAILY_DATA].cell(row=linha_errada, column=COL_DATA).value = INICIO + timedelta(days=10)
    wb.save(planilha)

    with pytest.raises(GradeDesalinhadaError):
        servico.adotar(corredor)

    repo = ActivityRepository(banco.connect())
    assert repo.por_periodo("p001", INICIO, cutover - timedelta(days=1)) == ()


def test_planilha_ausente_propaga_planilhanaoencontradaerror(servico, tmp_path):
    corredor = _corredor(tmp_path / "nao-existe.xlsx", cutover_date=INICIO + timedelta(days=2))

    with pytest.raises(PlanilhaNaoEncontradaError):
        servico.adotar(corredor)


def test_planilha_em_uso_propaga_planilhaemusoerror(planilha, servico, monkeypatch):
    import src.services.excel_service as modulo_excel_service

    def falhar(*_args, **_kwargs):
        raise PermissionError("arquivo em uso")

    monkeypatch.setattr(modulo_excel_service.openpyxl, "load_workbook", falhar)
    corredor = _corredor(planilha, cutover_date=INICIO + timedelta(days=2))

    with pytest.raises(PlanilhaEmUsoError):
        servico.adotar(corredor)


def test_duas_adocoes_com_conexoes_distintas_nao_duplicam(planilha, tmp_path):
    """Simula duas execuções concorrentes apontando para o mesmo arquivo de banco."""
    from src.services.excel_service import ExcelService

    cutover = INICIO + timedelta(days=2)
    corredor = _corredor(planilha, cutover_date=cutover)

    wb = openpyxl.load_workbook(planilha)
    _preencher(wb, corredor, INICIO, 10.0, 50.0)
    wb.save(planilha)

    caminho_banco = tmp_path / "estado.db"
    banco1 = DatabaseService(caminho_banco)
    banco1.init_schema()
    banco2 = DatabaseService(caminho_banco)

    try:
        AdocaoService(ExcelService(), ActivityRepository(banco1.connect())).adotar(corredor)
        AdocaoService(ExcelService(), ActivityRepository(banco2.connect())).adotar(corredor)

        achadas = ActivityRepository(banco1.connect()).por_periodo("p001", INICIO, INICIO)
        assert len(achadas) == 1  # não duplicou
    finally:
        banco1.close()
        banco2.close()
