"""Testes do SyncService (Fase 6): a orquestração ponta a ponta de um corredor.

O que se verifica aqui é a costura entre peças já testadas em separado —
`StravaClient`, `auth`, `ActivityService`, `ActivityRepository`,
`ExcelService`, `AdocaoService` —, não a lógica interna de cada uma.
`httpx.MockTransport` substitui o Strava; a planilha é uma cópia do template
real, como nos testes de Excel.
"""
from __future__ import annotations

import shutil
from datetime import UTC, date, datetime, timedelta
from pathlib import Path

import httpx
import openpyxl
import pytest

from src.api.rate_limiter import RateLimiter
from src.api.strava_client import StravaClient
from src.models.activity import Activity
from src.models.corredor import Corredor
from src.models.strava_token import StravaToken
from src.repositories.activity_repository import ActivityRepository
from src.repositories.corredor_state_repository import CorredorStateRepository
from src.repositories.escrita_excel_repository import EscritaExcelRepository
from src.services.activity_service import ActivityService
from src.services.adocao_service import AdocaoService
from src.services.database_service import DatabaseService
from src.services.excel_service import ABA_DAILY_DATA, ExcelService
from src.services.sync_service import SyncService
from src.utils.errors import QuotaExhaustedError, RevokedTokenError

TEMPLATE = Path(__file__).resolve().parents[1] / "Cópia de Planilha_carga_corrida.xlsx"
INICIO = date(2026, 1, 1)
URL_API = "https://api.exemplo/v3"
URL_OAUTH = "https://oauth.exemplo/token"

pytestmark = pytest.mark.skipif(
    not TEMPLATE.is_file(), reason="template real não encontrado no repositório"
)


class _Servidor:
    def __init__(self, *respostas: httpx.Response) -> None:
        self.respostas = list(respostas)
        self.requisicoes: list[httpx.Request] = []

    def __call__(self, request: httpx.Request) -> httpx.Response:
        self.requisicoes.append(request)
        if not self.respostas:
            raise AssertionError(f"requisição inesperada: {request.method} {request.url}")
        return self.respostas.pop(0)


def _resposta(status: int = 200, *, json=None) -> httpx.Response:
    cabecalhos = {"X-RateLimit-Limit": "100,1000", "X-RateLimit-Usage": "1,1"}
    return httpx.Response(status, json=json if json is not None else [], headers=cabecalhos)


def _erro_strava(status: int, recurso: str) -> httpx.Response:
    return httpx.Response(status, json={"errors": [{"resource": recurso, "code": "invalid"}]})


def _corrida(identificador: int, dia: str, *, tipo: str = "Run", km: float = 10.0, minutos: float = 50.0) -> dict:
    return {
        "id": identificador,
        "name": f"Corrida {identificador}",
        "type": tipo,
        "start_date": f"{dia}T09:00:00Z",
        "start_date_local": f"{dia}T09:00:00Z",
        "distance": km * 1000,
        "moving_time": int(minutos * 60),
        "elapsed_time": int(minutos * 60) + 30,
    }


@pytest.fixture
def banco(tmp_path):
    servico = DatabaseService(tmp_path / "estado.db")
    servico.init_schema()
    try:
        yield servico
    finally:
        servico.close()


@pytest.fixture
def planilha(tmp_path) -> Path:
    destino = tmp_path / "p001.xlsx"
    shutil.copy(TEMPLATE, destino)
    return destino


def _corredor(planilha_path: Path, *, cutover_date: date = INICIO) -> Corredor:
    return Corredor(
        id="p001",
        nome="Ana",
        refresh_token="token-do-toml",
        start_date=INICIO,
        cutover_date=cutover_date,
        excel_path=planilha_path,
    )


def _cliente(servidor: _Servidor, *, limiter: RateLimiter | None = None) -> StravaClient:
    return StravaClient(
        "12345",
        "segredo",
        limiter=limiter or RateLimiter(pausa_s=0, dormir=lambda _: None),
        http=httpx.Client(transport=httpx.MockTransport(servidor)),
        url_api=URL_API,
        url_oauth=URL_OAUTH,
    )


def _servico(banco, servidor: _Servidor, *, hoje: date, limiter: RateLimiter | None = None) -> SyncService:
    conn = banco.connect()
    activity_repo = ActivityRepository(conn)
    excel_service = ExcelService()
    return SyncService(
        client=_cliente(servidor, limiter=limiter),
        activity_service=ActivityService(),
        activity_repository=activity_repo,
        corredor_state_repository=CorredorStateRepository(conn),
        excel_service=excel_service,
        escrita_excel_repository=EscritaExcelRepository(conn),
        adocao_service=AdocaoService(excel_service, activity_repo),
        hoje=lambda: hoje,
    )


def _com_token_valido(banco, corredor_id: str = "p001") -> None:
    """Semeia um access token já válido — o teste não passa por `/oauth/token`."""
    CorredorStateRepository(banco.connect()).salvar_token(
        corredor_id,
        StravaToken(
            access_token="access-valido",
            refresh_token="refresh-valido",
            expira_em=datetime.now(UTC) + timedelta(hours=6),
            athlete_id=777,
        ),
    )


def _daily_data(planilha: Path):
    return openpyxl.load_workbook(planilha)[ABA_DAILY_DATA]


# --------------------------------------------------------------- caminho feliz


def test_salva_atividades_novas_no_banco(banco, planilha):
    corredor = _corredor(planilha)
    _com_token_valido(banco)
    servidor = _Servidor(_resposta(json=[_corrida(1, "2026-01-01")]))

    resultado = _servico(banco, servidor, hoje=INICIO).sincronizar(corredor)

    assert resultado.atividades_novas == 1
    achadas = ActivityRepository(banco.connect()).por_dia("p001", INICIO)
    assert len(achadas) == 1
    assert achadas[0].id == 1


def test_grava_carga_no_excel_para_o_dia_com_corrida(banco, planilha):
    corredor = _corredor(planilha)
    _com_token_valido(banco)
    servidor = _Servidor(_resposta(json=[_corrida(1, "2026-01-01", km=12.5)]))

    _servico(banco, servidor, hoje=INICIO).sincronizar(corredor)

    assert _daily_data(planilha)["C2"].value == 12.5


def test_preenche_dias_de_descanso_ate_hoje(banco, planilha):
    """A grade é contígua — dia sem corrida (inclusive hoje) precisa de Carga 0."""
    corredor = _corredor(planilha)
    _com_token_valido(banco)
    servidor = _Servidor(_resposta(json=[_corrida(1, "2026-01-01")]))

    _servico(banco, servidor, hoje=date(2026, 1, 3)).sincronizar(corredor)

    aba = _daily_data(planilha)
    assert aba["C2"].value == 10.0  # dia 1: corrida
    assert aba["C3"].value == 0.0  # dia 2: descanso, preenchido
    assert aba["C4"].value == 0.0  # dia 3 (hoje): descanso, preenchido
    assert aba.cell(row=5, column=3).value is None  # dia 4: ainda não chegou


def test_filtra_tipos_que_nao_sao_corrida(banco, planilha):
    corredor = _corredor(planilha)
    _com_token_valido(banco)
    servidor = _Servidor(_resposta(json=[_corrida(1, "2026-01-01", tipo="Ride")]))

    resultado = _servico(banco, servidor, hoje=INICIO).sincronizar(corredor)

    assert resultado.atividades_novas == 0
    assert _daily_data(planilha)["C2"].value == 0.0


def test_atividade_malformada_e_descartada_sem_derrubar_o_lote(banco, planilha, caplog):
    corredor = _corredor(planilha)
    _com_token_valido(banco)
    corrida_ruim = _corrida(2, "2026-01-02")
    del corrida_ruim["distance"]  # campo obrigatório ausente
    servidor = _Servidor(_resposta(json=[_corrida(1, "2026-01-01"), corrida_ruim]))

    with caplog.at_level("WARNING"):
        resultado = _servico(banco, servidor, hoje=date(2026, 1, 2)).sincronizar(corredor)

    assert resultado.atividades_novas == 1
    assert "descartada" in caplog.text


def test_segunda_sincronizacao_atualiza_em_vez_de_duplicar(banco, planilha):
    corredor = _corredor(planilha)
    _com_token_valido(banco)

    servidor1 = _Servidor(_resposta(json=[_corrida(1, "2026-01-01", km=10.0)]))
    _servico(banco, servidor1, hoje=INICIO).sincronizar(corredor)

    servidor2 = _Servidor(_resposta(json=[_corrida(1, "2026-01-01", km=12.0)]))
    resultado = _servico(banco, servidor2, hoje=INICIO).sincronizar(corredor)

    assert resultado.atividades_novas == 0
    assert resultado.atividades_atualizadas == 1
    achadas = ActivityRepository(banco.connect()).por_dia("p001", INICIO)
    assert len(achadas) == 1
    assert achadas[0].distance_m == 12000.0
    assert _daily_data(planilha)["C2"].value == 12.0


def test_dia_antigo_e_reescrito_quando_atividade_muda_de_data(banco, planilha):
    """`dias_afetados` alcança dias fora da janela recente — não só 'hoje'."""
    corredor = _corredor(planilha)
    _com_token_valido(banco)

    servidor1 = _Servidor(
        _resposta(json=[_corrida(1, "2026-01-01", km=10.0), _corrida(2, "2026-01-05", km=5.0)])
    )
    _servico(banco, servidor1, hoje=date(2026, 1, 5)).sincronizar(corredor)

    # Atividade 1 muda de 01/01 para 01/03 no Strava; "hoje" avança só um dia.
    servidor2 = _Servidor(
        _resposta(json=[_corrida(1, "2026-01-03", km=10.0), _corrida(2, "2026-01-05", km=5.0)])
    )
    _servico(banco, servidor2, hoje=date(2026, 1, 6)).sincronizar(corredor)

    aba = _daily_data(planilha)
    assert aba["C2"].value == 0.0  # 01/01: a corrida saiu de lá, carga volta a 0
    assert aba["C4"].value == 10.0  # 01/03: chegou aqui, fora da janela "recente"


# ------------------------------------------------------------------- after=


def test_after_usa_meia_noite_utc_do_cutover_na_primeira_vez(banco, planilha):
    corredor = _corredor(planilha, cutover_date=date(2026, 1, 5))
    _com_token_valido(banco)
    servidor = _Servidor(_resposta(json=[]))

    _servico(banco, servidor, hoje=date(2026, 1, 5)).sincronizar(corredor)

    pedido = servidor.requisicoes[0]
    esperado = int(datetime(2026, 1, 5, tzinfo=UTC).timestamp())
    assert int(pedido.url.params["after"]) == esperado


def test_after_usa_ultimo_evento_menos_a_folga_depois_da_primeira(banco, planilha):
    corredor = _corredor(planilha)
    _com_token_valido(banco)

    servidor1 = _Servidor(_resposta(json=[_corrida(1, "2026-01-10")]))
    _servico(banco, servidor1, hoje=date(2026, 1, 10)).sincronizar(corredor)

    servidor2 = _Servidor(_resposta(json=[]))
    _servico(banco, servidor2, hoje=date(2026, 1, 11)).sincronizar(corredor)

    pedido = servidor2.requisicoes[0]
    esperado = int((datetime(2026, 1, 10, 9, 0, tzinfo=UTC) - timedelta(days=2)).timestamp())
    assert int(pedido.url.params["after"]) == esperado


# ---------------------------------------------------------- atividade apagada


def test_detecta_atividade_apagada_no_strava(banco, planilha, caplog):
    corredor = _corredor(planilha)
    _com_token_valido(banco)
    ActivityRepository(banco.connect()).salvar(
        "p001",
        Activity(
            id=999,
            name="Corrida sumida",
            date=datetime(2026, 1, 5, 8, 0),  # noqa: DTZ001 — local e ingênuo, como o modelo
            distance_m=8000.0,
            moving_time_s=2400,
            elapsed_time_s=2450,
            start_date_utc=datetime(2026, 1, 5, 11, 0, tzinfo=UTC),
        ),
    )
    servidor = _Servidor(_resposta(json=[]))  # Strava não devolve mais nada

    with caplog.at_level("WARNING"):
        resultado = _servico(banco, servidor, hoje=date(2026, 1, 6)).sincronizar(corredor)

    assert resultado.atividades_apagadas_no_strava == 1
    assert "999" in caplog.text
    # Não apaga: carga continua no banco e reflete na planilha.
    assert len(ActivityRepository(banco.connect()).por_dia("p001", date(2026, 1, 5))) == 1
    assert _daily_data(planilha).cell(row=6, column=3).value == 8.0


# -------------------------------------------------------------------- adoção


def test_roda_adocao_quando_ha_periodo_manual(banco, planilha):
    cutover = date(2026, 1, 3)
    corredor = _corredor(planilha, cutover_date=cutover)
    _com_token_valido(banco)

    wb = openpyxl.load_workbook(planilha)
    wb[ABA_DAILY_DATA]["B2"] = INICIO
    wb[ABA_DAILY_DATA]["C2"] = 7.5
    wb.save(planilha)

    servidor = _Servidor(_resposta(json=[]))
    resultado = _servico(banco, servidor, hoje=cutover).sincronizar(corredor)

    assert resultado.adocao is not None
    assert resultado.adocao.dias_importados == 1


def test_nao_roda_adocao_quando_corredor_e_novo(banco, planilha):
    corredor = _corredor(planilha)  # start == cutover
    _com_token_valido(banco)
    servidor = _Servidor(_resposta(json=[]))

    resultado = _servico(banco, servidor, hoje=INICIO).sincronizar(corredor)

    assert resultado.adocao is None


# ------------------------------------------------------------- edição humana


def test_edicao_humana_depois_do_corte_e_preservada(banco, planilha):
    corredor = _corredor(planilha)
    _com_token_valido(banco)

    servidor1 = _Servidor(_resposta(json=[_corrida(1, "2026-01-01", km=10.0, minutos=50.0)]))
    _servico(banco, servidor1, hoje=INICIO).sincronizar(corredor)

    wb = openpyxl.load_workbook(planilha)
    wb[ABA_DAILY_DATA]["C2"] = 42.0  # humano corrige na mão
    wb.save(planilha)

    servidor2 = _Servidor(_resposta(json=[_corrida(1, "2026-01-01", km=15.0, minutos=60.0)]))
    _servico(banco, servidor2, hoje=INICIO).sincronizar(corredor)

    assert _daily_data(planilha)["C2"].value == 42.0


# -------------------------------------------------------------------- falhas


def test_revoked_token_propaga_sem_gravar_nada(banco, planilha):
    corredor = _corredor(planilha)
    servidor = _Servidor(_erro_strava(400, "RefreshToken"))

    with pytest.raises(RevokedTokenError):
        _servico(banco, servidor, hoje=INICIO).sincronizar(corredor)

    assert ActivityRepository(banco.connect()).por_dia("p001", INICIO) == ()
    estado = CorredorStateRepository(banco.connect()).buscar("p001")
    assert estado.precisa_reinscricao is True


def test_quota_exhausted_propaga_sem_gastar_requisicao(banco, planilha):
    corredor = _corredor(planilha)
    _com_token_valido(banco)

    limiter = RateLimiter(pausa_s=0, dormir=lambda _: None)
    limiter.registrar_resposta({"X-RateLimit-Limit": "100,1000", "X-RateLimit-Usage": "100,1000"})
    servidor = _Servidor()  # não deve ser chamado

    with pytest.raises(QuotaExhaustedError):
        _servico(banco, servidor, hoje=INICIO, limiter=limiter).sincronizar(corredor)

    assert servidor.requisicoes == []
